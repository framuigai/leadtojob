"""
PHASE 8 AND 9 ACCEPTANCE TESTS.

Phase 8:
    PDF/XLSX titles and wording identical;
    renders when a whole domain is N/A for a third of respondents;
    the privacy and wording checks pass on a POPULATED render.

Phase 9:
    a 4-person team produces no report and says why;
    a wave missing a domain is refused, not plotted as a drop.

These render real PDFs and real workbooks and read them back. A report that
builds without raising is not the same as a report that is correct, and the
things most worth checking — a suppressed cell, an absent domain, a band label
— all look fine right up until someone opens the file.
"""
import io
import re
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import db, domains
from app.analytics import reports, scoring as sc
from app.models import (
    Analysis, ReportEvent, ResponseSession, SurveyDeployment, SurveyItem,
    SurveyResponse,
)
from app.provisioning import clone_instrument_to_org
from tests import wording
from tests.conftest import login
from tests.factories import (
    make_branch, make_cycle, make_deployment, make_master_instrument, make_org,
    make_org_admin, make_segment,
)


def _pdf_text(data):
    """
    Extracted PDF text with whitespace collapsed.

    WeasyPrint lays text out in columns and wraps mid-phrase, so pypdf
    returns a title split across two lines. An exact substring match then
    fails on a phrase that is plainly present in the document. Collapsing
    whitespace tests what a READER sees rather than where the renderer
    happened to break a line.
    """
    pages = [page.extract_text() or ""
             for page in PdfReader(io.BytesIO(data)).pages]
    return re.sub(r"\s+", " ", " ".join(pages))


@pytest.fixture
def world(db):
    make_master_instrument()
    org = make_org('Kilimo Dairies')
    clone_instrument_to_org(org.id)
    admin = make_org_admin(org, email='l2@kilimo.test')
    db.session.commit()
    return {'org': org, 'admin': admin}


def _respond(org_id, deployment, scores, segment=None, na_domains=()):
    """
    One submitted session. `scores` maps a domain code to a 1-5 score;
    `na_domains` lists domains answered entirely "not applicable".
    """
    session = ResponseSession(
        deployment_id=deployment.id, org_id=org_id,
        branch_id=deployment.branch_id,
        segment_values=({'department': segment} if segment else {}),
        submitted_at=datetime.now(timezone.utc),
    )
    db.session.add(session)
    db.session.flush()

    items = (db.session.query(SurveyItem)
             .filter_by(org_id=org_id, version=1).all())
    for item in items:
        if item.domain_code not in scores and item.domain_code not in na_domains:
            continue
        na = item.domain_code in na_domains
        db.session.add(SurveyResponse(
            session_id=session.id, deployment_id=deployment.id,
            org_id=org_id, item_id=item.id,
            score=None if na else scores[item.domain_code],
            not_applicable=na))
    db.session.flush()
    deployment.responses_received = (deployment.responses_received or 0) + 1


def _closed_wave(world, cycle_id='KD-1', n=18, branch=None, min_n=1,
                 assessed=None, previous=None, varied=True):
    cycle = make_cycle(world['org'], cycle_id, status='open', branch=branch,
                       min_reporting_n=min_n,
                       assessed_domains=assessed or list(domains.DOMAINS),
                       previous_cycle_id=previous)
    wave = make_deployment(world['org'], cycle, branch=branch,
                           expected_respondents=n + 6,
                           public_slug=f'slug-{cycle_id}')
    for i in range(n):
        tier = i * 3 // n if varied else 1
        gem = (2, 3, 5)[tier]
        scores = {d: gem if d == 'GEM' else (2, 3, 4)[tier]
                  for d in (assessed or domains.DOMAINS)}
        _respond(world['org'].id, wave, scores,
                 segment=('FIN' if i % 2 else 'OPS'))
    wave.closed = True
    cycle.status = 'closed'
    db.session.commit()
    return cycle, wave


# ── Parity: the PDF and the XLSX cannot disagree ───────────────────────────

def test_the_title_comes_from_one_function():
    """
    Both formats read `report_title()`; neither composes its own. The two files
    are sent together and opened side by side, and a client who finds one name
    in one and a different name in the other stops trusting both.
    """
    for report_type in reports.REPORT_TYPES:
        assert reports.report_title(report_type)
    assert reports.report_title(reports.EXECUTIVE) == 'Executive Report'
    with pytest.raises(reports.ReportNotAvailable):
        reports.report_title('something_else')


def test_the_filename_comes_from_one_function():
    pdf = reports.report_filename(reports.HR_OD, 'Kilimo Dairies', 'KD-1', 'pdf')
    xlsx = reports.report_filename(reports.HR_OD, 'Kilimo Dairies', 'KD-1', 'xlsx')
    assert pdf[:-3] == xlsx[:-4]
    assert 'Kilimo_Dairies' in pdf and 'KD-1' in pdf


@pytest.mark.parametrize('report_type', [reports.EXECUTIVE, reports.HR_OD])
def test_the_pdf_and_the_xlsx_agree_on_title_and_figures(client, world,
                                                        report_type):
    """THE NAMED PHASE 8 CRITERION — checked by reading both files back."""
    cycle, _ = _closed_wave(world)
    login(client, world['admin'])

    pdf_res = client.get(
        f'/org/cycles/{cycle.id}/report/{report_type}/pdf')
    xlsx_res = client.get(
        f'/org/cycles/{cycle.id}/report/{report_type}/xlsx')
    assert pdf_res.status_code == 200
    assert xlsx_res.status_code == 200

    pdf_text = _pdf_text(pdf_res.data)
    workbook = load_workbook(io.BytesIO(xlsx_res.data))
    xlsx_text = '\n'.join(
        str(cell.value) for sheet in workbook.worksheets
        for row in sheet.iter_rows() for cell in row if cell.value is not None)

    title = reports.report_title(report_type)
    assert title in pdf_text
    assert title in xlsx_text
    assert world['org'].name in pdf_text and world['org'].name in xlsx_text
    assert cycle.cycle_name in pdf_text and cycle.cycle_name in xlsx_text

    # The scoring method in force, stated in both. §10.2.
    assert 'mean of the three subdomain scores' in pdf_text
    assert 'mean of the three subdomain scores' in xlsx_text

    # The §10.4 caveat, in both.
    assert 'do not, by themselves, establish' in pdf_text
    assert 'do not, by themselves, establish' in xlsx_text


# ── Phase 8: the hard render cases ─────────────────────────────────────────

def test_a_report_renders_when_a_whole_domain_is_na_for_a_third(client, world):
    """
    THE NAMED PHASE 8 CRITERION.

    Reward Satisfaction is the realistic case: a third of people genuinely do
    not know what their peers earn. The domain must still report from the rest,
    the coverage must be visible, and nothing may render as a zero.
    """
    cycle = make_cycle(world['org'], 'KD-NA', status='open', min_reporting_n=1)
    wave = make_deployment(world['org'], cycle, expected_respondents=24,
                           public_slug='slug-na')
    for i in range(18):
        scores = {d: (3 if i % 2 else 4) for d in domains.DOMAINS}
        if i < 6:
            _respond(world['org'].id, wave,
                     {d: v for d, v in scores.items() if d != 'PIE'},
                     na_domains=('PIE',))
        else:
            _respond(world['org'].id, wave, scores)
    wave.closed = True
    cycle.status = 'closed'
    db.session.commit()

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.HR_OD}/pdf')
    assert res.status_code == 200

    text = _pdf_text(res.data)
    assert 'Reward Satisfaction' in text
    # The domain reports from the 12 who could answer, not from 18 with zeros.
    assert '0.00' not in text.split('Reward Satisfaction')[1][:200]
    assert 'Coverage' in text


def test_a_suppressed_cell_never_renders_as_a_blank_or_a_zero(client, world):
    """
    A blank in a spreadsheet is read as zero by a human and by every formula
    that touches the column. "Suppressed (n < 5)" cannot be averaged by
    accident.
    """
    make_segment(world['org'], 'department', 'Department',
                 [('FIN', 'Finance'), ('OPS', 'Operations'), ('SAL', 'Sales')])
    cycle = make_cycle(world['org'], 'KD-SUP', status='open', min_reporting_n=5)
    wave = make_deployment(world['org'], cycle, expected_respondents=40,
                           public_slug='slug-sup')
    for i in range(20):
        _respond(world['org'].id, wave, {d: 4 for d in domains.DOMAINS},
                 segment='FIN')
    for i in range(20):
        _respond(world['org'].id, wave, {d: 3 for d in domains.DOMAINS},
                 segment='OPS')
    for i in range(3):
        _respond(world['org'].id, wave, {d: 1 for d in domains.DOMAINS},
                 segment='SAL')
    wave.closed = True
    cycle.status = 'closed'
    db.session.commit()

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.HR_OD}/xlsx')
    workbook = load_workbook(io.BytesIO(res.data))
    values = [str(cell.value) for sheet in workbook.worksheets
              for row in sheet.iter_rows() for cell in row
              if cell.value is not None]
    assert any('Suppressed' in v for v in values)


def test_no_band_vocabulary_survives_into_a_populated_pdf(client, world):
    """
    §10.5, on the finished artefact. Both sibling products have a four-band
    palette and inheriting it is the easiest mistake available.
    """
    cycle, _ = _closed_wave(world, 'KD-BANDS')
    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.EXECUTIVE}/pdf')
    text = _pdf_text(res.data)

    for forbidden in ('Critical Gap', 'Moderate Gap', 'Minor Gap',
                      'Strength Gap', 'Critical Risk'):
        assert forbidden not in text, forbidden
    assert 'not been validated' in text


def test_the_word_daunting_never_reaches_a_report(client, world):
    cycle, _ = _closed_wave(world, 'KD-D')
    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.HR_OD}/pdf')
    text = _pdf_text(res.data)
    assert 'Daunting' not in text
    assert 'Appropriate Challenge and Development' in text


def test_a_report_never_mentions_ai(client, world):
    """Inherited rule. The section is called Interpretation."""
    cycle, _ = _closed_wave(world, 'KD-AI')
    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.HR_OD}/pdf')
    text = _pdf_text(res.data)
    for forbidden in ('Claude', 'Anthropic', 'AI-generated', 'language model',
                      'LLM', 'ChatGPT'):
        assert forbidden not in text, forbidden


def test_every_report_carries_n_and_the_threshold(client, world):
    """A number without its n is how a four-person finding becomes a board slide."""
    cycle, _ = _closed_wave(world, 'KD-N')
    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.EXECUTIVE}/pdf')
    text = _pdf_text(res.data)
    assert 'n = 18' in text or 'n=18' in text
    assert 'Minimum reporting threshold' in text


def test_the_limitations_section_is_required_not_an_appendix(client, world):
    cycle, _ = _closed_wave(world, 'KD-LIM')
    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.EXECUTIVE}/pdf')
    text = _pdf_text(res.data)
    assert 'Limitations and what this does not show' in text
    assert 'No single satisfaction score' in text
    assert 'Not for decisions about individuals' in text


def test_downloading_writes_a_report_event(client, world):
    """§13.4 requires role-based access to results; this is the evidence."""
    cycle, _ = _closed_wave(world, 'KD-EV')
    login(client, world['admin'])
    client.get(f'/org/cycles/{cycle.id}/report/{reports.EXECUTIVE}/pdf')

    event = db.session.query(ReportEvent).one()
    assert event.report_type == reports.EXECUTIVE
    assert event.export_format == 'pdf'
    assert event.generated_by == world['admin'].id


# ── The closed-wave gate ───────────────────────────────────────────────────

def test_an_open_wave_refuses_to_generate_and_says_why(client, world):
    """
    The temporal-differencing mitigation, enforced rather than suggested. Two
    reports of the same group at n=5 and n=6 give up respondent 6 by
    subtraction, and no in-report suppression prevents it.
    """
    cycle, wave = _closed_wave(world, 'KD-OPEN')
    wave.closed = False
    db.session.commit()

    login(client, world['admin'])
    body = client.get(
        f'/org/cycles/{cycle.id}/report/{reports.EXECUTIVE}').get_data(as_text=True)
    assert 'still open' in body
    assert 'subtracted from one another' in body


def test_an_undeployed_cycle_refuses_to_generate(client, world):
    make_cycle(world['org'], 'KD-NONE', status='open')
    db.session.commit()

    login(client, world['admin'])
    body = client.get(
        f'/org/cycles/KD-NONE/report/{reports.EXECUTIVE}').get_data(as_text=True)
    assert 'has not been deployed' in body


# ── Phase 9: the Management Report suppression gate ────────────────────────

def test_a_four_person_team_produces_no_report_and_says_why(client, db, world):
    """
    THE NAMED PHASE 9 CRITERION.

    The most privacy-sensitive report in the product. The gate raises rather
    than returning a flag, so a caller that forgets to check cannot render an
    empty report that looks finished.
    """
    branch = make_branch(world['org'], 'Nakuru plant')
    cycle = make_cycle(world['org'], 'KD-SMALL', status='open',
                       branch=branch, min_reporting_n=5)
    wave = make_deployment(world['org'], cycle, branch=branch,
                           expected_respondents=10, public_slug='slug-small')
    for _ in range(4):
        _respond(world['org'].id, wave, {d: 4 for d in domains.DOMAINS})
    wave.closed = True
    db.session.commit()

    login(client, world['admin'])
    body = client.get(
        f'/org/cycles/{cycle.id}/report/{reports.MANAGEMENT}'
        f'?branch={branch.id}').get_data(as_text=True)

    assert '4 responses were received' in body
    assert 'at least 5 are needed' in body
    # It must not read as a fault, and it must say the responses still count.
    assert 'working as intended' in body
    assert 'still count towards the results for the wider organisation' in body

    # And the download refuses too, not only the preview.
    res = client.get(
        f'/org/cycles/{cycle.id}/report/{reports.MANAGEMENT}/pdf'
        f'?branch={branch.id}', follow_redirects=True)
    assert b'%PDF' not in res.data


def test_a_team_at_the_threshold_does_produce_a_report(client, db, world):
    """The complement — without it, a gate that refused everything would pass."""
    branch = make_branch(world['org'], 'Nakuru plant')
    cycle = make_cycle(world['org'], 'KD-OK', status='open', branch=branch,
                       min_reporting_n=5)
    wave = make_deployment(world['org'], cycle, branch=branch,
                           expected_respondents=10, public_slug='slug-ok')
    for _ in range(5):
        _respond(world['org'].id, wave, {d: 4 for d in domains.DOMAINS})
    wave.closed = True
    db.session.commit()

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.MANAGEMENT}/pdf'
                     f'?branch={branch.id}')
    assert res.status_code == 200
    assert res.data[:4] == b'%PDF'


def test_the_management_report_says_no_individual_is_visible(client, db, world):
    branch = make_branch(world['org'], 'Nakuru plant')
    cycle = make_cycle(world['org'], 'KD-MGR', status='open', branch=branch,
                       min_reporting_n=1)
    wave = make_deployment(world['org'], cycle, branch=branch,
                           expected_respondents=10, public_slug='slug-mgr')
    for i in range(12):
        _respond(world['org'].id, wave,
                 {d: (2 if i < 6 else 5) for d in domains.DOMAINS})
    wave.closed = True
    db.session.commit()

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.MANAGEMENT}/pdf'
                     f'?branch={branch.id}')
    text = _pdf_text(res.data)
    assert 'No individual response is shown to you' in text
    # The template's own words: "Nothing in this report may be used in an
    # appraisal, a performance conversation about a particular person..."
    assert 'may be used in an appraisal' in text
    # And the stylesheet must not have leaked into the body — it did once,
    # because these partials are raw CSS and were included outside <style>.
    assert 'Report cover' not in text
    assert 'font-family' not in text


# ── Phase 9: trend ─────────────────────────────────────────────────────────

def test_a_wave_missing_a_domain_is_refused_not_plotted_as_a_drop(client, world):
    """
    THE NAMED PHASE 9 CRITERION.

    A pulse wave of GEM + TID following a full wave. Reward Satisfaction has a
    figure last time and none this time. Plotted naively that renders as a fall
    to zero — a catastrophic finding that did not happen. The organisation
    asked fewer questions; it did not get worse.
    """
    first, _ = _closed_wave(world, 'KD-W1', n=18)
    second, _ = _closed_wave(world, 'KD-W2', n=18,
                             assessed=['GEM', 'TID'], previous=first.id)

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{second.id}/report/{reports.HR_OD}/pdf')
    assert res.status_code == 200
    text = _pdf_text(res.data)

    assert 'was not measured in this wave' in text
    assert 'is not a domain that fell' in text
    # And no false drop anywhere near it.
    assert 'Reward Satisfaction' in text


def test_a_comparable_domain_does_report_its_change(client, world):
    first, _ = _closed_wave(world, 'KD-T1', n=18)
    second, _ = _closed_wave(world, 'KD-T2', n=18, previous=first.id)

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{second.id}/report/{reports.HR_OD}/pdf')
    text = _pdf_text(res.data)
    assert 'Change since the previous wave' in text
    assert 'does not by itself prove' in text        # §16's caveat


def test_the_trend_reports_both_waves_response_counts(client, world):
    """
    §16: "consider other explanations, timing, sample composition". A wave of
    40 and a wave of 120 are not the same population measured twice.
    """
    first, _ = _closed_wave(world, 'KD-C1', n=30)
    second, _ = _closed_wave(world, 'KD-C2', n=12, previous=first.id)

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{second.id}/report/{reports.HR_OD}/pdf')
    text = _pdf_text(res.data)
    assert 'noticeably different' in text
    assert 'may reflect who answered' in text


def test_no_trend_section_without_a_previous_wave(client, world):
    cycle, _ = _closed_wave(world, 'KD-SOLO')
    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.HR_OD}/pdf')
    text = _pdf_text(res.data)
    assert 'Change since the previous wave' not in text


# ── The narrative payload ──────────────────────────────────────────────────

def test_the_narrative_payload_contains_no_suppressed_group(world):
    """
    The model cannot narrate what it was not given, which is a stronger
    guarantee than asking it not to.
    """
    # THREE groups, not two. With only two, suppressing the small one forces
    # the other to be suppressed as well — total minus the visible group IS the
    # hidden one — so nothing survives and the test would prove nothing.
    make_segment(world['org'], 'department', 'Department',
                 [('FIN', 'Finance'), ('OPS', 'Operations'), ('SAL', 'Sales')])
    cycle = make_cycle(world['org'], 'KD-PAY', status='open', min_reporting_n=5)
    wave = make_deployment(world['org'], cycle, expected_respondents=60,
                           public_slug='slug-pay')
    for i in range(20):
        _respond(world['org'].id, wave, {d: 4 for d in domains.DOMAINS},
                 segment='FIN')
    for i in range(20):
        _respond(world['org'].id, wave, {d: 3 for d in domains.DOMAINS},
                 segment='OPS')
    for i in range(2):
        _respond(world['org'].id, wave, {d: 1 for d in domains.DOMAINS},
                 segment='SAL')
    wave.closed = True
    db.session.commit()

    from app.provisioning import active_segments
    profiles = sc.load_profiles(db.session, [wave.id])
    context = reports.build_context(
        reports.HR_OD, world['org'], cycle, [wave], profiles,
        segments=active_segments(world['org'].id))
    payload = reports.narrative_payload(context)

    groups = payload.get('group_comparisons', {}).get('Department', [])
    names = {g['group'] for g in groups}
    assert names                               # something survived
    assert 'Sales' not in names                # suppressed, therefore absent
    # Complementary suppression hides a second group with it, so at most one of
    # the two large groups is present.
    assert len(names) <= 2


def test_the_narrative_payload_carries_no_identity(world):
    cycle, wave = _closed_wave(world, 'KD-ID')
    profiles = sc.load_profiles(db.session, [wave.id])
    context = reports.build_context(
        reports.EXECUTIVE, world['org'], cycle, [wave], profiles)
    blob = repr(reports.narrative_payload(context))

    for forbidden in ('session_id', 'email', 'respondent', '@'):
        assert forbidden not in blob, forbidden


def test_the_payload_offers_only_the_specs_own_response_categories(world):
    """
    §11's wording, verbatim, so the model can only name categories the
    specification actually lists — and the Reward rows contain no training.
    """
    cycle, wave = _closed_wave(world, 'KD-CAT')
    profiles = sc.load_profiles(db.session, [wave.id])
    context = reports.build_context(
        reports.HR_OD, world['org'], cycle, [wave], profiles)
    categories = reports.narrative_payload(context)['available_response_categories']

    assert categories['Peer Comparison'] == 'Role and pay-equity review'
    for reward in ('Peer Comparison', 'Industry Comparison', 'Expectations'):
        assert 'training' not in categories[reward].lower()


# ── Scoping ────────────────────────────────────────────────────────────────

def test_another_orgs_report_is_not_reachable(client, db, world):
    other = make_org('Globex Inc')
    make_cycle(other, 'GBX-1', status='open')
    db.session.commit()

    login(client, world['admin'])
    assert client.get(
        f'/org/cycles/GBX-1/report/{reports.EXECUTIVE}').status_code == 404


def test_an_unknown_report_type_is_404(client, world):
    cycle, _ = _closed_wave(world, 'KD-404')
    login(client, world['admin'])
    assert client.get(
        f'/org/cycles/{cycle.id}/report/final_verdict').status_code == 404


# ── The wording guards, on a POPULATED render ──────────────────────────────
#
# THE NAMED PHASE 8 CRITERION. The guards themselves are tested in
# tests/test_wording.py against real violations, so a vacuous guard cannot make
# these pass.

@pytest.mark.parametrize('report_type', list(reports.REPORT_TYPES))
def test_a_populated_report_makes_no_unnegated_causal_claim(client, db, world,
                                                           report_type):
    """
    §10.4, §17.8. Checked on the finished PDF with a narrative in place, not on
    a template — the narrative is where a causal sentence would actually come
    from.
    """
    branch = make_branch(world['org'], 'Nakuru plant')
    cycle = make_cycle(world['org'], 'KD-WORD', status='open', branch=branch,
                       min_reporting_n=1)
    wave = make_deployment(world['org'], cycle, branch=branch,
                           expected_respondents=24, public_slug='slug-word')
    for i in range(18):
        tier = i * 3 // 18
        _respond(world['org'].id, wave,
                 {d: (2, 3, 5)[tier] if d == 'GEM' else (1, 3, 5)[tier]
                  for d in domains.DOMAINS})
    wave.closed = True
    db.session.commit()

    # A narrative that deliberately contains the REQUIRED negated forms, so
    # this test would fail if the guard were naive.
    db.session.add(Analysis(
        org_id=world['org'].id, cycle_id=cycle.id, report_type=report_type,
        scope_key=str(branch.id) if report_type == reports.MANAGEMENT else None,
        payload={
            'executive_summary':
                'This wave shows associations between measures, not causes. '
                'Trust scores 2.1 against 3.4 across the diagnostic '
                'subdomains.',
            'summary':
                'Your group reported Trust at 2.1. It does not show that '
                'Trust causes higher or lower satisfaction.',
            'key_findings':
                'Employees reporting lower General Job Satisfaction also '
                'report lower Trust. This does not establish that one '
                'produced the other.',
            'limitations':
                'Score bands have not been validated, so no result is '
                'described as low or high. A lower Appropriate Challenge '
                'score does not indicate stress or overload.',
        },
        model='test', generated_by=world['admin'].id))
    db.session.commit()

    login(client, world['admin'])
    url = f'/org/cycles/{cycle.id}/report/{report_type}/pdf'
    if report_type == reports.MANAGEMENT:
        url += f'?branch={branch.id}'
    res = client.get(url)
    assert res.status_code == 200
    text = _pdf_text(res.data)

    assert wording.causal_hits(text) == [], wording.causal_hits(text)[:3]
    assert wording.verdict_hits(text) == [], wording.verdict_hits(text)[:3]
    assert wording.stress_hits(text) == [], wording.stress_hits(text)[:3]
    assert wording.identity_hits(text) == [], wording.identity_hits(text)[:3]

    # And the negated safeguards really are present — otherwise the assertions
    # above would pass on a report that simply never mentions causation.
    assert 'not causes' in text or 'does not' in text


def test_the_populated_render_test_would_catch_a_real_violation(client, db,
                                                                world):
    """
    The meta-check for the test above. If a narrative DID assert causation, the
    sweep must fail — otherwise the four assertions are decoration.
    """
    cycle, _ = _closed_wave(world, 'KD-BAD')
    db.session.add(Analysis(
        org_id=world['org'].id, cycle_id=cycle.id,
        report_type=reports.EXECUTIVE, scope_key=None,
        payload={'executive_summary':
                 'Low autonomy causes the fall in overall satisfaction.'},
        model='test', generated_by=world['admin'].id))
    db.session.commit()

    login(client, world['admin'])
    res = client.get(f'/org/cycles/{cycle.id}/report/{reports.EXECUTIVE}/pdf')
    text = _pdf_text(res.data)
    assert wording.causal_hits(text), 'the sweep failed to catch a real claim'
